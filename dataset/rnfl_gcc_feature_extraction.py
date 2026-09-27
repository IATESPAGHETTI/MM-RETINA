"""
Structured feature extraction for the RNFL/GCC workbook
(dataset/hvf_and)rnfl_gcc/RNFL and GCC data/Glaucoma AI.xlsx).

This is the exploratory-only clinical dataset: 171 eye-level rows with no
patient ID, no date, no name (see dataset/rnfl_gcc_provenance_audit.md —
a forensic pass over the workbook's hidden sheets, formulas, comments,
metadata, and defined names confirmed no identifier is recoverable). It
CANNOT be linked to the HVF DICOM cohort or to GAMMA, and it cannot be
patient-grouped for cross-validation because there is no patient key to
group by — any k-fold split of these 171 rows has an unknown risk of
putting fellow eyes of the same real (but unidentifiable) patient into
different folds. That limitation is carried through into every place this
CSV is consumed and must never be silently dropped.

Usage:
    python dataset/rnfl_gcc_feature_extraction.py --out dataset/rnfl_gcc_features.csv
"""

from __future__ import annotations

import argparse
import os

import openpyxl
import pandas as pd

XLSX_PATH = os.path.join(
    os.path.dirname(__file__), "hvf_and)rnfl_gcc", "RNFL and GCC data", "Glaucoma AI.xlsx"
)

SEVERITY_FROM_LABEL = {"MILD": 0, "MODERATE": 1, "SEVERE": 2}

# Exactly the fields MM_RETINA_MULTIDATASET_REPRESENTATION_AND_CASE_EXAMPLES.md
# specifies for the RNFL/GCC branch: Average RNFL, RNFL regional
# measurements, Average GCC, GCC regional measurements, age, gender.
NUMERIC_FEATURES = [
    "Average_RNFL", "RNFL_Superior", "RNFL_Inferior", "RNFL_Nasal", "RNFL_Temporal",
    "Average_GCC", "GCC_Superior", "GCC_Inferior",
    "GCC_Supero_nasal", "GCC_Supero_temporal", "GCC_Infero_nasal", "GCC_Infero_temporal",
    "age",
]
CATEGORICAL_FEATURES = ["gender"]  # eye is kept in the CSV for reference but not used as a feature


def extract(xlsx_path: str = XLSX_PATH) -> pd.DataFrame:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    header = [str(h).strip().replace(" ", "_") for h in header]

    records = []
    for r in data:
        rec = dict(zip(header, r))
        records.append({
            "si_no": rec["SI_NO"],
            "age": rec["Age"],
            "gender": str(rec["Gender"]).strip() if rec["Gender"] is not None else None,
            "eye": rec["Eye"],
            "severity_label_str": rec["Glaucoma_Severity"],
            "severity_label": SEVERITY_FROM_LABEL[str(rec["Glaucoma_Severity"]).strip().upper()],
            "Average_RNFL": rec["Average_RNFL"],
            "RNFL_Superior": rec["RNFL_Superior"],
            "RNFL_Inferior": rec["RNFL_Inferior"],
            "RNFL_Nasal": rec["RNFL_Nasal"],
            "RNFL_Temporal": rec["RNFL_Temporal"],
            "Average_GCC": rec["Average_GCC"],
            "GCC_Superior": rec["GCC_Superior"],
            "GCC_Inferior": rec["GCC_Inferior"],
            "GCC_Supero_nasal": rec["GCC_Supero_nasal"],
            "GCC_Supero_temporal": rec["GCC_Supero_temporal"],
            "GCC_Infero_nasal": rec["GCC_Infero_nasal"],
            "GCC_Infero_temporal": rec["GCC_Infero_temporal"],
        })
    return pd.DataFrame(records)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", default=XLSX_PATH)
    ap.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "rnfl_gcc_features.csv"))
    args = ap.parse_args()

    df = extract(args.xlsx)
    df.to_csv(args.out, index=False)
    print(f"Extracted {len(df)} eye-level RNFL/GCC rows (no patient ID — see "
          f"dataset/rnfl_gcc_provenance_audit.md)")
    print(f"Severity distribution: {df['severity_label_str'].value_counts().to_dict()}")
    print(f"Wrote {args.out}")
